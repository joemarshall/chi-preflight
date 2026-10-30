import re
import zipfile

from docx import Document

from ..findings import Finding, Level

NAME = "docx-authors"
CATEGORY = "anonymity"
LEVEL = "quick"
ORCID_RE = re.compile(r"\b\d{4}-\d{4}-\d{4}-\d{3}[\dXx]\b")


def _xml_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels")}


def check(path, config):
    anon = config.get("anonymous", True)
    out = []
    doc = Document(path)
    cp = doc.core_properties
    parts = _xml_parts(path)
    if anon:
        for label, val in (("Author", cp.author), ("Last modified by", cp.last_modified_by)):
            if (val or "").strip():
                out.append(Finding(Level.ERROR, NAME, path, f"{label} metadata present: {val!r}"))
        app = parts.get("docProps/app.xml", "")
        for tag in ("Company", "Manager"):
            m = re.search(r"<%s>([^<]+)</%s>" % (tag, tag), app)
            if m:
                out.append(Finding(Level.ERROR, NAME, path, f"{tag} property present: {m.group(1)!r}"))
        names = set()
        for n, x in parts.items():
            if n.startswith("word/") and not n.endswith(".rels"):
                names |= set(re.findall(r'w:author="([^"]+)"', x))
        for n in sorted(names):
            out.append(Finding(Level.ERROR, NAME, path, f"Comment/tracked-change author present: {n!r}"))
        if ORCID_RE.search(" ".join(p.text for p in doc.paragraphs)) or "orcid.org" in " ".join(parts.values()):
            out.append(Finding(Level.ERROR, NAME, path, "ORCID found in document"))
        emails = {e for p in doc.paragraphs[:30] for e in re.findall(r"[\w.+-]+@[\w-]+\.[\w.-]+", p.text)}
        for e in sorted(emails):
            out.append(Finding(Level.ERROR, NAME, path, f"Email address in document: {e}"))
    else:
        author = (cp.author or "").strip()
        if not author:
            out.append(Finding(Level.ERROR, NAME, path, "Author metadata is missing"))
            return out
        n_auth = len([a for a in re.split(r"[;]|\band\b", author) if a.strip()])
        blob = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(parts.values())
        orcids = set(ORCID_RE.findall(blob))
        if not orcids:
            out.append(Finding(Level.WARNING, NAME, path, "No ORCID iDs found in document"))
        elif len(orcids) < n_auth:
            out.append(Finding(Level.WARNING, NAME, path, f"{len(orcids)} ORCID iD(s) found for {n_auth} author(s)"))
        out.append(Finding(Level.INFO, NAME, path, f"Author metadata: {author!r}; {len(orcids)} ORCID iD(s)"))
    return out

