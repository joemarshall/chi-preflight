from docx import Document

from ..findings import Finding, Level

NAME = "docx-structure"


def check(path, config):
    out = []
    doc = Document(path)
    paras = doc.paragraphs
    if not any(p.text.strip() for p in paras):
        out.append(Finding(Level.ERROR, NAME, path, "Document has no text"))
    headings = [p for p in paras if p.style.name.lower().startswith("heading")]
    if not headings:
        out.append(Finding(Level.WARNING, NAME, path, "No heading styles used"))
    empty = sum(1 for p in paras if not p.text.strip())
    if empty > 10:
        out.append(Finding(Level.WARNING, NAME, path, f"{empty} empty paragraphs (use spacing instead)"))
    shapes = doc.inline_shapes
    out.append(Finding(Level.INFO, NAME, path,
                       f"{len(paras)} paragraphs, {len(headings)} headings, {len(doc.tables)} tables, {len(shapes)} inline images"))
    return out

