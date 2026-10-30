import re

from docx import Document

from ..findings import Finding, Level

NAME = "docx-wordcount"
BIB_RE = re.compile(r"(?:\d+\.?\s*)?bibliography:?|(?:\d+\.?\s*)?references:?", re.I)
LABEL_RE = re.compile(r"(?:figure|fig\.|table|listing|algorithm|equation)\s*\d+", re.I)


def _is_heading(p):
    return p.style.name.lower().startswith(("heading", "title"))


def _is_label(p):
    return p.style.name.lower().startswith("caption") or bool(LABEL_RE.match(p.text.strip()))


def check(path, config):
    # doc.paragraphs covers body paragraphs only: table cells, headers and footers are excluded
    doc = Document(path)
    words = skipped_caps = skipped_bib = 0
    in_bib = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if BIB_RE.fullmatch(text) and _is_heading(p):
            in_bib = True
            continue
        if in_bib and _is_heading(p):
            in_bib = False
        n = len(text.split())
        if in_bib:
            skipped_bib += 1
        elif _is_label(p):
            skipped_caps += 1
        else:
            words += n
    out = [Finding(Level.INFO, NAME, path,
                   f"Word count: {words} (excluding bibliography [{skipped_bib} entries], "
                   f"captions/labels [{skipped_caps}], tables, headers and footers)")]
    limit = config.get("max_words")
    if limit and words > limit:
        out.append(Finding(Level.ERROR, NAME, path, f"Word count {words} exceeds maximum of {limit}"))
    return out
