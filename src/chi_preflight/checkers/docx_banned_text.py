from docx import Document

from ..findings import Finding, Level
from ._banned import find_banned

NAME = "docx-banned-text"
CATEGORY = "content"


def _paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from part.paragraphs


def check(path, config):
    banned = config.get("banned_text")
    if not banned:
        return []
    doc = Document(path)
    text = "\n".join(p.text for p in _paragraphs(doc))
    return [Finding(Level.ERROR, NAME, path, f"{message} (found: '{phrase}')")
            for phrase, message in find_banned(text, banned)]
